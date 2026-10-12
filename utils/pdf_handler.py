
from pathlib import Path

import PyPDF2
from docx import Document


def extract_pages_from_pdf(file):
    """Extract text from each PDF page and retain page numbers."""

    reader = PyPDF2.PdfReader(file)
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()

        if text:
            pages.append({
                "page": page_number,
                "text": text,
            })

    return pages


def extract_pages_from_docx(file):
    """
    Extract paragraphs and table text from a DOCX file.

    DOCX files do not have reliable fixed page boundaries because
    pagination depends on the Word layout and rendering environment.
    We use section numbers instead of inventing page numbers.
    """

    document = Document(file)
    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # Include text stored inside Word tables.
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip().replace("\n", " ")
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                paragraphs.append(" | ".join(cells))

    if not paragraphs:
        return []

    # Divide text into numbered sections for source identification.
    full_text = "\n".join(paragraphs)
    words = full_text.split()
    chunk_size = 180
    overlap = 40
    step = chunk_size - overlap

    sections = []

    for start in range(0, len(words), step):
        text = " ".join(words[start:start + chunk_size]).strip()

        if text:
            sections.append({
                "page": None,
                "section": len(sections) + 1,
                "text": text,
            })

        if start + chunk_size >= len(words):
            break

    return sections


def extract_document(file, filename):
    """Extract text from a supported PDF or DOCX file."""

    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        return extract_pages_from_pdf(file)

    if extension == ".docx":
        return extract_pages_from_docx(file)

    raise ValueError("Unsupported file type. Upload a PDF or DOCX file.")


def extract_text_from_pdf(file):
    """Maintain compatibility with existing PDF extraction calls."""

    pages = extract_pages_from_pdf(file)

    return "\n\n".join(
        item["text"] for item in pages
    )
